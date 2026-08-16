#!/usr/bin/env python3
from __future__ import annotations
import hashlib, json, re, sys, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

ROOT = Path(__file__).resolve().parent

def sha(path):
    h=hashlib.sha256()
    with path.open('rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
    return h.hexdigest()

def fail(msg):
    print('FAIL:',msg); raise SystemExit(1)

def main():
    checks=[]
    registry=ROOT/'canonical'/'AF_v6_2_CANONICAL_REGISTRY.json'
    schema=ROOT/'canonical'/'AF_v6_2_case_record.schema.json'
    tests=ROOT/'canonical'/'AF_v6_2_MIXED_STATE_TEST_VECTORS.json'
    for p in [registry,schema,tests]:
        if not p.exists(): fail(f'missing {p.relative_to(ROOT)}')
    reg=json.loads(registry.read_text(encoding='utf-8'))
    if reg.get('package_version')!='6.2': fail('registry package version')
    for key in ['decision_status','evidence_state','review_state']:
        if key not in reg.get('three_axis_record',{}): fail(f'missing axis {key}')
    if 'non_erasure_invariant' not in reg: fail('missing non-erasure invariant')
    checks.append('canonical registry')

    docx=list(ROOT.rglob('*.docx'))
    pdf=list(ROOT.rglob('*.pdf'))
    xlsx=list(ROOT.rglob('*.xlsx'))
    if not docx: fail('no DOCX files')
    if not pdf: fail('no PDF files')
    for p in docx:
        with zipfile.ZipFile(p) as z:
            if z.testzip(): fail(f'bad DOCX zip {p}')
            if 'word/document.xml' not in z.namelist(): fail(f'not DOCX {p}')
            ET.fromstring(z.read('word/document.xml'))
    checks.append(f'{len(docx)} DOCX containers')
    for p in pdf:
        if p.read_bytes()[:5]!=b'%PDF-': fail(f'bad PDF header {p}')
    checks.append(f'{len(pdf)} PDF containers')
    for p in xlsx:
        with zipfile.ZipFile(p) as z:
            if z.testzip(): fail(f'bad XLSX zip {p}')
            ET.fromstring(z.read('xl/workbook.xml'))
        if load_workbook:
            wb=load_workbook(p,read_only=False,data_only=False)
            if 'v6.2 Controls' not in wb.sheetnames or 'v6.2 Test Vectors' not in wb.sheetnames:
                fail(f'v6.2 workbook sheets missing in {p}')
    checks.append(f'{len(xlsx)} workbook containers')

    core_docs=[]
    for p in ROOT.rglob('*.docx'):
        if 'core_protocol' in p.name.lower(): core_docs.append(p)
    if not core_docs: fail('core protocol DOCX not found')
    found=False
    for p in core_docs:
        with zipfile.ZipFile(p) as z:
            txt=' '.join(ET.fromstring(z.read('word/document.xml')).itertext())
            if 'Release 6.2 control clarifications' in txt and 'non-erasure' in txt.lower(): found=True
    if not found: fail('core v6.2 clarifications missing')
    checks.append('core semantic controls')

    inv=ROOT/'SHA256SUMS.txt'
    if not inv.exists(): fail('SHA256SUMS.txt missing')
    for line in inv.read_text(encoding='utf-8').splitlines():
        if not line.strip(): continue
        digest, rel=line.split('  ',1)
        target=ROOT/rel
        if not target.exists(): fail(f'hash target missing {rel}')
        if sha(target)!=digest: fail(f'hash mismatch {rel}')
    checks.append('hash inventory')

    print('PASS: Auditable Flourishing v6.2')
    for c in checks: print(' -',c)

if __name__=='__main__': main()
